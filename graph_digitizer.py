import math
import os
import re
import subprocess
import sys

import numpy as np

try:
    from openpyxl import Workbook
    from openpyxl.chart import Reference, ScatterChart, Series
    from openpyxl.chart.axis import ChartLines
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.chart.layout import Layout, ManualLayout
except ImportError:  # pragma: no cover
    Workbook = None
    Reference = None
    ScatterChart = None
    Series = None
    ChartLines = None
    GraphicalProperties = None
    Layout = None
    ManualLayout = None


def _project_python_path():
    candidates = [
        os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "..", "Scripts", "python.exe")),
        r"C:\Python314\python.exe",
        r"C:\Python3.14\python.exe",
        sys.executable,
    ]
    for candidate in candidates:
        if candidate and os.path.isfile(candidate):
            return candidate
    return candidates[0]


def _is_usable_python(candidate):
    if not candidate or not os.path.isfile(candidate):
        return False
    try:
        subprocess.run(
            [candidate, "-c", "import sys; print(sys.executable)"],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            check=True,
            timeout=10,
        )
    except (OSError, subprocess.SubprocessError, ValueError):
        return False
    return True


def _restart_with_project_interpreter():
    if __name__ != "__main__":
        return

    project_python = _project_python_path()
    target_python = project_python if _is_usable_python(project_python) else None
    if target_python is None:
        return

    current_python = os.path.normcase(os.path.realpath(sys.executable))
    target_python_real = os.path.normcase(os.path.realpath(target_python))
    if current_python != target_python_real:
        os.execv(target_python, [target_python, os.path.abspath(__file__), *sys.argv[1:]])


_restart_with_project_interpreter()

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

try:
    import fitz  # type: ignore[import-not-found]
except ImportError:  # pragma: no cover
    fitz = None

try:
    from PIL import Image, ImageTk  # type: ignore[import-not-found]
except ImportError:  # pragma: no cover
    Image = None
    ImageTk = None


def _pdf_number(text):
    try:
        return float(text.strip().replace(",", ""))
    except ValueError:
        return None


def _axis_limits_from_pdf_ticks(ticks, lower, upper):
    ticks = sorted(ticks, key=lambda item: item[0])
    if len(ticks) < 2 or ticks[0][1] <= 0 or ticks[-1][1] <= 0:
        return None

    first_position, first_value = ticks[0]
    last_position, last_value = ticks[-1]
    if first_position == last_position or first_value == last_value:
        return None

    first_log = math.log10(first_value)
    slope = (math.log10(last_value) - first_log) / (last_position - first_position)

    def value_at(position):
        nearest = min(ticks, key=lambda tick: abs(tick[0] - position))
        if abs(nearest[0] - position) <= 2.5:
            return nearest[1]
        estimate = 10 ** (first_log + slope * (position - first_position))
        exponent = math.floor(math.log10(estimate))
        nice_values = [
            multiplier * 10**power
            for power in (exponent - 1, exponent, exponent + 1)
            for multiplier in (1, 2, 5)
        ]
        nearest_nice_value = min(nice_values, key=lambda value: abs(value - estimate))
        if abs(nearest_nice_value - estimate) / estimate <= 0.05:
            return nearest_nice_value
        return estimate

    lower_value = value_at(lower)
    upper_value = value_at(upper)
    return (lower_value, upper_value) if lower_value < upper_value else (upper_value, lower_value)


def _extract_pdf_curve_names(page, drawings, curve_drawings):
    text_spans = []
    for block in page.get_text("dict").get("blocks", ()):
        for line in block.get("lines", ()):
            for span in line.get("spans", ()):
                x0, y0, x1, y1 = span["bbox"]
                text = span.get("text", "").strip()
                if text:
                    text_spans.append(((x0 + x1) / 2, (y0 + y1) / 2, x1, text))

    curve_left = min(drawing["rect"].x0 for drawing in curve_drawings)
    rating_swatches = []
    for drawing in drawings:
        bounds = drawing["rect"]
        fill = drawing.get("fill")
        if (
            fill is None
            or bounds.x1 >= curve_left
            or not 5 <= bounds.width <= 30
            or not 2 <= bounds.height <= 15
        ):
            continue

        center_y = (bounds.y0 + bounds.y1) / 2
        row_ratings = [
            (abs(y - center_y), value)
            for _, y, x1, text in text_spans
            if x1 <= bounds.x0
            and abs(y - center_y) <= 2.5
            and (value := _pdf_number(text)) is not None
            and value > 0
        ]
        if not row_ratings:
            continue

        _, rating = min(row_ratings, key=lambda item: item[0])
        rating_text = str(int(rating)) if rating.is_integer() else format(rating, "g")
        rating_swatches.append((fill, f"{rating_text}A"))

    if len(rating_swatches) != len(curve_drawings):
        return None

    color_matches = []
    for curve_index, curve in enumerate(curve_drawings):
        curve_color = curve.get("color")
        if curve_color is None:
            return None
        for swatch_index, (swatch_color, _) in enumerate(rating_swatches):
            distance = math.sqrt(
                sum((curve_channel - swatch_channel) ** 2
                    for curve_channel, swatch_channel in zip(curve_color, swatch_color))
            )
            color_matches.append((distance, curve_index, swatch_index))

    names = [None] * len(curve_drawings)
    used_swatches = set()
    for distance, curve_index, swatch_index in sorted(color_matches):
        if distance > 0.45:
            break
        if names[curve_index] is not None or swatch_index in used_swatches:
            continue
        names[curve_index] = rating_swatches[swatch_index][1]
        used_swatches.add(swatch_index)

    return names if all(names) else None


def _detect_pdf_graphs(pdf_path):
    if fitz is None:
        return []

    graph_titles = ("characteristic", "derating curve", "graph")
    detected_graphs = []
    with fitz.open(pdf_path) as document:
        for page_index, page in enumerate(document):
            panels = []
            for drawing in page.get_drawings():
                bounds = drawing["rect"]
                items = drawing.get("items", ())
                if (
                    drawing.get("fill") is None
                    or len(items) != 1
                    or items[0][0] != "re"
                    or bounds.width < 100
                    or bounds.height < 80
                    or not 0.45 <= bounds.width / bounds.height <= 2.5
                ):
                    continue
                panels.append(bounds)

            page_graphs = []
            for block in page.get_text("dict").get("blocks", ()):
                for line in block.get("lines", ()):
                    for span in line.get("spans", ()):
                        title = span.get("text", "").strip()
                        title_lower = title.lower()
                        if not title or not any(term in title_lower for term in graph_titles):
                            continue

                        title_x0, _, title_x1, title_y1 = span["bbox"]
                        panel_candidates = []
                        for panel in panels:
                            vertical_gap = panel.y0 - title_y1
                            horizontal_overlap = max(
                                0,
                                min(panel.x1, title_x1) - max(panel.x0, title_x0),
                            )
                            if (
                                0 <= vertical_gap <= 80
                                and horizontal_overlap >= min(30, (title_x1 - title_x0) * 0.25)
                            ):
                                panel_candidates.append(
                                    (vertical_gap, -panel.width * panel.height, panel)
                                )
                        if not panel_candidates:
                            continue

                        _, _, panel = min(panel_candidates, key=lambda item: (item[0], item[1]))
                        if any(graph["plot_bounds"] == tuple(panel) for graph in page_graphs):
                            continue
                        page_graphs.append(
                            {
                                "title": title,
                                "page_index": page_index,
                                "plot_bounds": tuple(panel),
                                "page_size": (page.rect.width, page.rect.height),
                            }
                        )
            detected_graphs.extend(page_graphs)
    return detected_graphs


def _extract_pdf_curves(pdf_path, page_index=0, requested_plot_bounds=None):
    if fitz is None:
        return None

    with fitz.open(pdf_path) as document:
        if not document.page_count:
            return None
        if not 0 <= page_index < document.page_count:
            return None
        page = document[page_index]
        page_width = page.rect.width
        page_height = page.rect.height
        drawings = page.get_drawings()

        curve_drawings = []
        for drawing in drawings:
            bounds = drawing["rect"]
            color = drawing.get("color")
            items = drawing.get("items", ())
            if (
                color is None
                or drawing.get("width") is None
                or drawing["width"] < 0.7
                or len(items) < 3
                or not all(item[0] == "c" for item in items)
                or bounds.width < 20
                or bounds.height < 100
                or bounds.height / bounds.width < 2
            ):
                continue
            curve_drawings.append(drawing)

        if not curve_drawings:
            return None

        curve_drawings.sort(key=lambda drawing: min(item[1].x for item in drawing["items"]))
        curve_names = _extract_pdf_curve_names(page, drawings, curve_drawings)
        frame_candidates = []
        for drawing in drawings:
            color = drawing.get("color")
            items = drawing.get("items", ())
            bounds = drawing["rect"]
            if (
                color is None
                or max(color) > 0.15
                or drawing.get("width") is None
                or drawing["width"] < 0.05
                or len(items) != 2
                or not all(item[0] == "l" for item in items)
                or bounds.width < 80
                or bounds.height < 80
                or not 0.3 <= bounds.width / bounds.height <= 2
            ):
                continue

            enclosed = sum(
                bounds.x0 <= drawing_curve["rect"].x0
                and drawing_curve["rect"].x1 <= bounds.x1
                and drawing_curve["rect"].y1 > bounds.y0
                and drawing_curve["rect"].y0 < bounds.y1
                for drawing_curve in curve_drawings
            )
            if enclosed >= 3:
                frame_candidates.append((enclosed, bounds.width * bounds.height, bounds))

        if requested_plot_bounds is None:
            if not frame_candidates:
                return None
            _, _, plot_bounds = max(frame_candidates, key=lambda candidate: (candidate[0], -candidate[1]))
        else:
            plot_bounds = fitz.Rect(requested_plot_bounds)
            curve_drawings = [
                drawing
                for drawing in curve_drawings
                if drawing["rect"].x1 >= plot_bounds.x0
                and drawing["rect"].x0 <= plot_bounds.x1
                and drawing["rect"].y1 >= plot_bounds.y0
                and drawing["rect"].y0 <= plot_bounds.y1
            ]
            if len(curve_drawings) < 2:
                return None
            curve_drawings.sort(key=lambda drawing: min(item[1].x for item in drawing["items"]))
            curve_names = _extract_pdf_curve_names(page, drawings, curve_drawings)

        x_ticks = []
        y_ticks = []
        x_axis_title = None
        y_axis_title = None
        graph_title = None
        for block in page.get_text("dict").get("blocks", ()):
            for line in block.get("lines", ()):
                for span in line.get("spans", ()):
                    x0, y0, x1, y1 = span["bbox"]
                    center_y = (y0 + y1) / 2
                    center_x = (x0 + x1) / 2
                    text = span.get("text", "").strip()
                    direction = line.get("dir", (1.0, 0.0))
                    if not text:
                        continue
                    if (
                        graph_title is None
                        and plot_bounds.x0 - 20 <= center_x <= plot_bounds.x1 + 20
                        and plot_bounds.y0 - 28 <= y1 <= plot_bounds.y0
                    ):
                        graph_title = text
                    if (
                        x_axis_title is None
                        and plot_bounds.x0 <= center_x <= plot_bounds.x1
                        and plot_bounds.y1 <= y0 <= plot_bounds.y1 + 35
                    ):
                        value = _pdf_number(text)
                        if value is None:
                            x_axis_title = text
                    if (
                        y_axis_title is None
                        and abs(direction[0]) < 0.2
                        and abs(direction[1]) > 0.8
                        and plot_bounds.x0 - 35 <= x1 <= plot_bounds.x0
                        and plot_bounds.y0 <= center_y <= plot_bounds.y1
                    ):
                        value = _pdf_number(text)
                        if value is None:
                            y_axis_title = text

                    value = _pdf_number(text)
                    if value is None or value <= 0:
                        continue
                    if plot_bounds.y1 <= y0 <= plot_bounds.y1 + 18 and plot_bounds.x0 - 2 <= center_x <= plot_bounds.x1 + 2:
                        x_ticks.append((center_x, value))
                    if plot_bounds.x0 - 42 <= x1 <= plot_bounds.x0 + 2 and plot_bounds.y0 - 2 <= center_y <= plot_bounds.y1 + 2:
                        y_ticks.append((center_y, value))

        x_limits = _axis_limits_from_pdf_ticks(x_ticks, plot_bounds.x0, plot_bounds.x1)
        y_limits = _axis_limits_from_pdf_ticks(y_ticks, plot_bounds.y0, plot_bounds.y1)
        if x_limits is None or y_limits is None:
            return None

        curves = []
        for index, drawing in enumerate(curve_drawings, start=1):
            sampled_points = []
            pixel_points = []
            for item in drawing["items"]:
                _, start, control1, control2, end = item
                for step in range(41):
                    if sampled_points and step == 0:
                        continue
                    t = step / 40
                    inverse_t = 1 - t
                    x = (
                        inverse_t**3 * start.x
                        + 3 * inverse_t**2 * t * control1.x
                        + 3 * inverse_t * t**2 * control2.x
                        + t**3 * end.x
                    )
                    y = (
                        inverse_t**3 * start.y
                        + 3 * inverse_t**2 * t * control1.y
                        + 3 * inverse_t * t**2 * control2.y
                        + t**3 * end.y
                    )
                    sampled_points.append((x, y))
                    if plot_bounds.x0 <= x <= plot_bounds.x1 and plot_bounds.y0 <= y <= plot_bounds.y1:
                        pixel_points.append((x, y))

            top_intersections = []
            for first, second in zip(sampled_points, sampled_points[1:]):
                first_x, first_y = first
                second_x, second_y = second
                if (first_y - plot_bounds.y0) * (second_y - plot_bounds.y0) > 0:
                    continue
                if first_y == second_y:
                    continue
                ratio = (plot_bounds.y0 - first_y) / (second_y - first_y)
                x = first_x + ratio * (second_x - first_x)
                if plot_bounds.x0 <= x <= plot_bounds.x1:
                    top_intersections.append((x, plot_bounds.y0))

            if top_intersections:
                pixel_points.append(min(top_intersections, key=lambda point: point[0]))
            elif len(sampled_points) >= 2:
                upper_points = sorted(sampled_points, key=lambda point: point[1])[:2]
                (first_x, first_y), (second_x, second_y) = sorted(
                    upper_points, key=lambda point: point[1]
                )
                if first_y != second_y:
                    x = first_x + (
                        (plot_bounds.y0 - first_y)
                        * (second_x - first_x)
                        / (second_y - first_y)
                    )
                    if plot_bounds.x0 <= x <= plot_bounds.x1:
                        pixel_points.append((x, plot_bounds.y0))

            pixel_points.sort(key=lambda point: (point[0], point[1]))
            points = []
            for x, y in pixel_points:
                if points and x <= points[-1][0]:
                    continue
                points.append((x / page_width, y / page_height))
            if len(points) >= 2:
                color = drawing["color"]
                curves.append(
                    {
                        "name": curve_names[index - 1] if curve_names else f"Curve {index}",
                        "color": "#{:02x}{:02x}{:02x}".format(
                            *(round(channel * 255) for channel in color)
                        ),
                        "points": points,
                    }
                )

    if len(curves) < 2:
        return None
    return {
        "curves": curves,
        "plot_bounds": (
            plot_bounds.x0 / page_width,
            plot_bounds.y0 / page_height,
            plot_bounds.x1 / page_width,
            plot_bounds.y1 / page_height,
        ),
        "axis_values": {
            "xMin": x_limits[0],
            "xMax": x_limits[1],
            "yMin": y_limits[0],
            "yMax": y_limits[1],
        },
        "axis_scales": {"x": "log", "y": "log"},
        "chart_title": graph_title or "Digitized graph curves",
        "axis_titles": {"x": x_axis_title or "X", "y": y_axis_title or "Y"},
    }


def _extract_pdf_legend_entries(page, plot_bounds):
    text_spans = []
    for block in page.get_text("dict").get("blocks", ()):
        for line in block.get("lines", ()):
            for span in line.get("spans", ()):
                text = span.get("text", "").strip()
                if text:
                    text_spans.append((span["bbox"], text))

    entries = []
    for drawing in page.get_drawings():
        bounds = drawing["rect"]
        items = drawing.get("items", ())
        color = drawing.get("color")
        if (
            color is None
            or len(items) != 1
            or items[0][0] != "l"
            or bounds.x0 <= plot_bounds.x1
            or not 10 <= bounds.width <= 20
            or bounds.height > 1.5
            or not plot_bounds.y0 - 5 <= bounds.y0 <= plot_bounds.y1 + 2
        ):
            continue

        center_y = (bounds.y0 + bounds.y1) / 2
        nearby_labels = []
        for text_bounds, text in text_spans:
            x0, y0, _, y1 = text_bounds
            match = re.fullmatch(r"(\d+(?:\.\d+)?)\s*A", text, re.IGNORECASE)
            if (
                match
                and bounds.x1 - 1 <= x0 <= bounds.x1 + 30
                and abs((y0 + y1) / 2 - center_y) <= 1.5
            ):
                nearby_labels.append((abs((y0 + y1) / 2 - center_y), match.group(1)))
        if not nearby_labels:
            continue

        _, rating = min(nearby_labels)
        entries.append(
            (
                bounds.x0,
                center_y,
                f"{format(float(rating), 'g')}A",
                tuple(round(channel * 255) for channel in color),
            )
        )

    entries.sort(key=lambda entry: (entry[0], entry[1]))
    names = [entry[2] for entry in entries]
    if len(entries) < 2 or len(set(names)) != len(names):
        return None
    return [
        {"name": name, "color_rgb": color}
        for _, _, name, color in entries
    ]


def _trace_raster_curve(plot_pixels, expected_color, occluding_colors=()):
    height, width, _ = plot_pixels.shape
    interior = plot_pixels[5:-5, 5:-5].astype(np.int16)
    expected = np.asarray(expected_color, dtype=np.int16)
    near_color = np.max(np.abs(interior - expected), axis=2) <= 15
    matching_pixels = interior[near_color]
    if not len(matching_pixels):
        return None

    colors, counts = np.unique(matching_pixels, axis=0, return_counts=True)
    trace_color = colors[int(np.argmax(counts))]
    color_mask = np.max(
        np.abs(plot_pixels.astype(np.int16) - trace_color),
        axis=2,
    ) <= 7
    color_mask[:, :5] = False
    color_mask[:, -5:] = False

    columns = []
    endpoint_candidates = []
    max_run_length = max(120, int(height * 0.2))
    for x in range(width):
        rows = np.flatnonzero(color_mask[:, x])
        if not len(rows):
            continue

        starts = np.r_[0, np.flatnonzero(np.diff(rows) > 1) + 1]
        ends = np.r_[starts[1:] - 1, len(rows) - 1]
        groups = []
        for start, end in zip(starts, ends):
            run_length = int(end - start + 1)
            if not 2 <= run_length < max_run_length:
                continue
            center_y = float((rows[start] + rows[end]) / 2)
            group = (center_y, run_length)
            groups.append(group)
            if 0.65 * height <= center_y <= 0.95 * height:
                endpoint_candidates.append((center_y, run_length, x))
        if groups:
            columns.append((x, groups))

    if not endpoint_candidates:
        return None

    _, _, endpoint_x = max(endpoint_candidates, key=lambda item: (item[0], item[1]))
    endpoint_groups = dict(columns)[endpoint_x]
    endpoint_y, _ = max(
        (
            group
            for group in endpoint_groups
            if 0.65 * height <= group[0] <= 0.95 * height
        ),
        key=lambda group: group[0],
    )
    traced_points = [(endpoint_x, endpoint_y)]
    previous_x, previous_y = endpoint_x, endpoint_y
    previous_slope = 0.0
    previous_point = None

    for x, groups in reversed([entry for entry in columns if entry[0] < endpoint_x]):
        predicted_y = previous_y + previous_slope * (x - previous_x)
        candidates = [group for group in groups if group[0] <= previous_y + 3]
        if not candidates:
            continue

        y, _ = min(
            candidates,
            key=lambda group: (
                abs(group[0] - predicted_y)
                + max(0, group[1] - 40) * 0.15
            ),
        )
        if previous_point is not None and x != previous_point[0]:
            previous_slope = max(
                0.0,
                min(30.0, (previous_y - previous_point[1]) / (previous_x - previous_point[0])),
            )
        previous_point = (previous_x, previous_y)
        traced_points.append((x, y))
        previous_x, previous_y = x, y

    traced_points.reverse()
    tail = np.asarray(traced_points[-8:], dtype=float)
    tail_slope, tail_intercept = np.polyfit(tail[:, 0], tail[:, 1], 1)
    extension_x, _ = traced_points[-1]
    for _ in range(max(1, int(width * 0.01))):
        next_x = extension_x + 1
        if next_x >= width:
            break
        rows = np.flatnonzero(color_mask[:, next_x])
        if not len(rows):
            break

        starts = np.r_[0, np.flatnonzero(np.diff(rows) > 1) + 1]
        ends = np.r_[starts[1:] - 1, len(rows) - 1]
        groups = [
            float((rows[start] + rows[end]) / 2)
            for start, end in zip(starts, ends)
        ]
        predicted_y = tail_slope * next_x + tail_intercept
        candidates = [
            y for y in groups
            if abs(y - predicted_y) <= max(2.0, abs(tail_slope) + 1.0)
        ]
        if not candidates:
            break
        extension_y = min(candidates, key=lambda y: abs(y - predicted_y))
        traced_points.append((next_x, extension_y))
        extension_x = next_x

    # Preserve terminal pixels hidden by an overlapping curve.
    for _ in range(2):
        next_x = extension_x + 1
        if (
            not occluding_colors
            or next_x >= width
            or np.any(color_mask[:, next_x])
        ):
            break

        predicted_y = tail_slope * next_x + tail_intercept
        other_color_mask = np.zeros(height, dtype=bool)
        pixels = plot_pixels[:, next_x, :].astype(np.int16)
        for color in occluding_colors:
            color_difference = np.max(
                np.abs(pixels - np.asarray(color, dtype=np.int16)),
                axis=1,
            )
            other_color_mask |= color_difference <= 7
        other_rows = np.flatnonzero(other_color_mask)
        if not len(other_rows) or np.min(np.abs(other_rows - predicted_y)) > 4:
            break

        traced_points.append((next_x, float(predicted_y)))
        extension_x = next_x

    if len(traced_points) < width * 0.15:
        return None

    x_positions = np.asarray([point[0] for point in traced_points], dtype=float)
    y_positions = np.maximum.accumulate(
        np.asarray([point[1] for point in traced_points], dtype=float)
    )
    if (
        x_positions[-1] - x_positions[0] < width * 0.1
        or y_positions[-1] - y_positions[0] < height * 0.7
    ):
        return None

    return list(zip(x_positions.tolist(), y_positions.tolist()))


def _extract_pdf_raster_curves(pdf_path, page_index, plot_bounds, chart_title):
    if fitz is None or Image is None:
        return None

    try:
        from rapidocr import RapidOCR
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError(
            "RapidOCR is required to digitize rasterized PDF charts. "
            "Install project dependencies with: python -m pip install -r requirements.txt"
        ) from exc

    with fitz.open(pdf_path) as document:
        if not 0 <= page_index < document.page_count:
            return None
        page = document[page_index]
        bounds = fitz.Rect(plot_bounds)
        legend_entries = _extract_pdf_legend_entries(page, bounds)
        if not legend_entries:
            return None

        clip = fitz.Rect(
            bounds.x0 - 35,
            bounds.y0 - 35,
            bounds.x1 + 45,
            bounds.y1 + 45,
        ) & page.rect
        ocr_pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False, clip=clip)
        ocr_image = Image.frombytes(
            "RGB",
            (ocr_pixmap.width, ocr_pixmap.height),
            ocr_pixmap.samples,
        )
        ocr_result = RapidOCR()(ocr_image)
        if (
            ocr_result is None
            or ocr_result.boxes is None
            or len(ocr_result.boxes) == 0
        ):
            return None

        plot_left = round(bounds.x0 * 2) - ocr_pixmap.x
        plot_right = round(bounds.x1 * 2) - ocr_pixmap.x
        plot_top = round(bounds.y0 * 2) - ocr_pixmap.y
        plot_bottom = round(bounds.y1 * 2) - ocr_pixmap.y
        x_ticks = []
        y_ticks = []
        x_axis_titles = []
        y_axis_titles = []

        for box, text, score in zip(ocr_result.boxes, ocr_result.txts, ocr_result.scores):
            if score < 0.35:
                continue
            center_x = sum(float(point[0]) for point in box) / len(box)
            center_y = sum(float(point[1]) for point in box) / len(box)
            value = _pdf_number(text)
            if value is not None and value > 0:
                if (
                    plot_left - 90 <= center_x <= plot_left - 2
                    and plot_top - 5 <= center_y <= plot_bottom + 5
                ):
                    y_ticks.append((center_y, value))
                if (
                    plot_left - 5 <= center_x <= plot_right + 5
                    and plot_bottom + 2 <= center_y <= plot_bottom + 55
                ):
                    x_ticks.append((center_x, value))
            elif value is None:
                if (
                    plot_left - 90 <= center_x <= plot_left - 2
                    and plot_top <= center_y <= plot_bottom
                ):
                    y_axis_titles.append(text.strip())
                if (
                    plot_left <= center_x <= plot_right
                    and plot_bottom < center_y <= plot_bottom + 60
                ):
                    x_axis_titles.append(text.strip())

        x_limits = _axis_limits_from_pdf_ticks(x_ticks, plot_left, plot_right)
        y_limits = _axis_limits_from_pdf_ticks(y_ticks, plot_top, plot_bottom)
        if x_limits is None or y_limits is None:
            return None

        curve_pixmap = page.get_pixmap(matrix=fitz.Matrix(4, 4), alpha=False, clip=clip)
        curve_image = Image.frombytes(
            "RGB",
            (curve_pixmap.width, curve_pixmap.height),
            curve_pixmap.samples,
        )
        image_pixels = np.asarray(curve_image)
        left = round(bounds.x0 * 4) - curve_pixmap.x
        right = round(bounds.x1 * 4) - curve_pixmap.x
        top = round(bounds.y0 * 4) - curve_pixmap.y
        bottom = round(bounds.y1 * 4) - curve_pixmap.y
        plot_pixels = image_pixels[top : bottom + 1, left : right + 1]
        if plot_pixels.size == 0:
            return None

        x_scale = 2 / ocr_pixmap.width
        y_scale = 2 / ocr_pixmap.height
        plot_origin_x = (bounds.x0 * 2 - ocr_pixmap.x) / ocr_pixmap.width
        plot_origin_y = (bounds.y0 * 2 - ocr_pixmap.y) / ocr_pixmap.height
        plot_width_ratio = (bounds.x1 - bounds.x0) * x_scale
        plot_height_ratio = (bounds.y1 - bounds.y0) * y_scale
        curves = []
        for entry in legend_entries:
            traced = _trace_raster_curve(
                plot_pixels,
                entry["color_rgb"],
                occluding_colors=[
                    other_entry["color_rgb"]
                    for other_entry in legend_entries
                    if other_entry is not entry
                ],
            )
            if not traced:
                return None

            if len(traced) < 2 or traced[1][1] <= traced[0][1]:
                return None
            first_points = traced[:8]
            first_y = np.asarray([point[1] for point in first_points], dtype=float)
            first_x = np.asarray([point[0] for point in first_points], dtype=float)
            if first_y[-1] > first_y[0]:
                start_x = float(np.polyfit(first_y, first_x, 1)[1])
            else:
                start_x = first_x[0]
            start_x = max(0.0, min(first_x[0], start_x))
            pixel_points = [(start_x, 0.0)] + [
                (x, y) for x, y in traced if y > 0
            ]

            points = []
            for x, y in pixel_points:
                normalized_x = float(
                    plot_origin_x
                    + plot_width_ratio * x / max(1, plot_pixels.shape[1] - 1)
                )
                normalized_y = float(
                    plot_origin_y
                    + plot_height_ratio * y / max(1, plot_pixels.shape[0] - 1)
                )
                if points and normalized_x <= points[-1][0]:
                    continue
                points.append((normalized_x, normalized_y))
            if len(points) < 2:
                return None

            red, green, blue = entry["color_rgb"]
            curves.append(
                {
                    "name": entry["name"],
                    "color": f"#{red:02x}{green:02x}{blue:02x}",
                    "points": points,
                }
            )

        return {
            "curves": curves,
            "plot_bounds": (
                plot_origin_x,
                plot_origin_y,
                plot_origin_x + plot_width_ratio,
                plot_origin_y + plot_height_ratio,
            ),
            "axis_values": {
                "xMin": x_limits[0],
                "xMax": x_limits[1],
                "yMin": y_limits[0],
                "yMax": y_limits[1],
            },
            "axis_scales": {"x": "log", "y": "log"},
            "chart_title": chart_title,
            "axis_titles": {
                "x": x_axis_titles[0] if x_axis_titles else "Current [A]",
                "y": y_axis_titles[0] if y_axis_titles else "Time [s]",
            },
        }


def _select_pdf_graph(root, graphs):
    if not graphs:
        return None
    if len(graphs) == 1:
        return graphs[0]

    dialog = tk.Toplevel(root)
    dialog.title("Select a graph")
    dialog.transient(root)
    dialog.resizable(False, False)
    dialog.grab_set()
    result = {"graph": None}

    tk.Label(
        dialog,
        text="Multiple graphs were detected. Select which one to process:",
        padx=16,
        pady=12,
        anchor="w",
    ).pack(fill="x")

    selected_index = tk.IntVar(value=0)
    for index, graph in enumerate(graphs):
        page_number = graph["page_index"] + 1
        tk.Radiobutton(
            dialog,
            text=f"Page {page_number}: {graph['title']}",
            variable=selected_index,
            value=index,
            anchor="w",
            padx=16,
        ).pack(fill="x")

    buttons = tk.Frame(dialog, padx=16, pady=12)
    buttons.pack(fill="x")

    def accept_selection():
        result["graph"] = graphs[selected_index.get()]
        dialog.destroy()

    ttk.Button(buttons, text="Process selected graph", command=accept_selection).pack(side="left")
    ttk.Button(buttons, text="Cancel", command=dialog.destroy).pack(side="right")
    dialog.protocol("WM_DELETE_WINDOW", dialog.destroy)
    root.wait_window(dialog)
    return result["graph"]


class GraphDigitizerApp:
    def __init__(self, root):
        self.root = root
        self.headless_mode = len(sys.argv) > 1
        self.root.title("Graph Digitizer")
        self.root.geometry("1180x760")
        self.root.minsize(900, 620)
        self.root.configure(bg="#f5f7fb")

        self.state = {
            "image": None,
            "image_path": None,
            "mode": None,
            "origin": None,
            "opposite": None,
            "points": [],
            "curves": [],
            "axis_scales": {"x": "linear", "y": "linear"},
            "axis_values": {"xMin": 0.0, "xMax": 10.0, "yMin": 0.0, "yMax": 10.0},
            "chart_title": "Digitized graph curves",
            "axis_titles": {"x": "X", "y": "Y"},
        }
        self.last_export_path = None
        self.zoom = 1.0
        self.display_photo = None
        self.display_size = (0, 0)
        self.range_drag_start = None
        self.range_drag_current = None

        self.style = ttk.Style()
        self.style.theme_use("clam")
        self.style.configure("TButton", padding=(10, 7), relief="flat", font=("Segoe UI", 10))
        self.style.configure("Accent.TButton", background="#365cf5", foreground="white")
        self.style.map("Accent.TButton", background=[("active", "#294bd8")])
        self.style.configure("TEntry", padding=6, font=("Segoe UI", 10))
        self.style.configure("Treeview", font=("Segoe UI", 10))
        self.style.configure("Treeview.Heading", font=("Segoe UI", 10, "bold"))

        self.number_format = self._make_number_format()

        self._build_ui()
        self._auto_load_from_cli()

    def _auto_load_from_cli(self):
        if len(sys.argv) <= 1:
            self.root.after(100, self.open_file_dialog)
            return
        if len(sys.argv) != 2:
            raise ValueError("Provide exactly one input path: an image or PDF file.")
        candidate = sys.argv[1]
        if not os.path.isfile(candidate):
            raise FileNotFoundError(f"Input file not found: {candidate}")
        supported_extensions = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff", ".pdf"}
        if os.path.splitext(candidate)[1].lower() not in supported_extensions:
            raise ValueError("Input must be a supported image file or PDF.")
        self.load_image(candidate)

    def _make_number_format(self):
        return lambda value: format(value, ".8g")

    def _build_ui(self):
        container = tk.Frame(self.root, bg="#f5f7fb", padx=28, pady=28)
        container.pack(fill="both", expand=True)

        tk.Label(
            container,
            text="Graph Digitizer",
            bg="#f5f7fb",
            fg="#19233a",
            font=("Segoe UI", 20, "bold"),
        ).pack(anchor="w")
        tk.Label(
            container,
            text="Enter or browse to an image/PDF. The graph is digitized and saved to Excel automatically.",
            bg="#f5f7fb",
            fg="#68738a",
            font=("Segoe UI", 10),
        ).pack(anchor="w", pady=(4, 18))

        path_row = tk.Frame(container, bg="#f5f7fb")
        path_row.pack(fill="x")
        tk.Label(
            path_row,
            text="Image/PDF location",
            bg="#f5f7fb",
            fg="#19233a",
            font=("Segoe UI", 10, "bold"),
        ).pack(anchor="w", pady=(0, 6))
        entry_row = tk.Frame(path_row, bg="#f5f7fb")
        entry_row.pack(fill="x")
        self.path_var = tk.StringVar()
        self.path_entry = ttk.Entry(entry_row, textvariable=self.path_var)
        self.path_entry.pack(side="left", fill="x", expand=True)
        self.path_entry.bind("<Return>", self._process_path_entry)
        ttk.Button(entry_row, text="Browse…", command=self.open_file_dialog).pack(side="left", padx=(8, 0))

        self.status_var = tk.StringVar(value="Choose a supported image or PDF file to begin.")
        tk.Label(
            container,
            textvariable=self.status_var,
            bg="#f5f7fb",
            fg="#59657c",
            font=("Segoe UI", 10),
            anchor="w",
            justify="left",
            wraplength=800,
        ).pack(fill="x", pady=(16, 0))

    def _build_control_sections(self, parent):
        sections = [
            ("Upload graph", "Choose an image of your chart to get started.", "upload"),
            ("Set axis scale", "Enter the values at the graph's edges.", "axis"),
            ("Calibrate graph", "Drag a rectangle around the graph's plotting area, inside the axes.", "calibration"),
            ("Digitize points", "Click the curve or chart to capture each data point.", "points"),
        ]

        for idx, (title, text, kind) in enumerate(sections, 1):
            section = tk.Frame(parent, bg="#ffffff", pady=0)
            section.pack(fill="x", pady=(0, 18))
            if idx < len(sections):
                section.config(padx=0)
            title_row = tk.Frame(section, bg="#ffffff")
            title_row.pack(fill="x")
            num = tk.Label(title_row, text=str(idx), bg="#eef1ff", fg="#365cf5", width=3, font=("Segoe UI", 10, "bold"), relief="flat")
            num.pack(side="left", padx=(0, 10))
            tk.Label(title_row, text=title, bg="#ffffff", fg="#19233a", font=("Segoe UI", 12, "bold")).pack(side="left")
            tk.Label(section, text=text, bg="#ffffff", fg="#68738a", font=("Segoe UI", 10), justify="left", anchor="w").pack(anchor="w", pady=(6, 10))

            if kind == "upload":
                ttk.Button(section, text="Choose image", command=self.open_file_dialog).pack(anchor="w")
                self.file_input = tk.StringVar()
                self.upload_button = None
            elif kind == "axis":
                axis_grid = tk.Frame(section, bg="#ffffff")
                axis_grid.pack(fill="x")
                self.axis_fields = {
                    "xMin": self._field(axis_grid, "X minimum"),
                    "xMax": self._field(axis_grid, "X maximum"),
                    "yMin": self._field(axis_grid, "Y minimum"),
                    "yMax": self._field(axis_grid, "Y maximum"),
                }
                self.axis_fields["xMin"].insert(0, "0")
                self.axis_fields["xMax"].insert(0, "10")
                self.axis_fields["yMin"].insert(0, "0")
                self.axis_fields["yMax"].insert(0, "10")
                for field in self.axis_fields.values():
                    field.bind("<KeyRelease>", lambda _event: self._on_axis_input())
            elif kind == "calibration":
                self.range_button = ttk.Button(section, text="Drag to select graph range", command=self.set_range_mode, state="disabled")
                self.range_button.pack(fill="x")
                self.message_var = tk.StringVar(value="")
                self.message_label = tk.Label(section, textvariable=self.message_var, bg="#ffffff", fg="#b43f42", font=("Segoe UI", 10), justify="left", wraplength=300)
                self.message_label.pack(anchor="w", pady=(8, 0))
            elif kind == "points":
                actions = tk.Frame(section, bg="#ffffff")
                actions.pack(fill="x")
                self.undo_button = ttk.Button(actions, text="Undo last", command=self.undo_last_point, state="disabled")
                self.undo_button.pack(side="left", expand=True, padx=(0, 8))
                self.clear_button = ttk.Button(actions, text="Clear points", command=self.clear_points, state="disabled")
                self.clear_button.pack(side="left", expand=True, padx=(0, 8))
                self.export_button = ttk.Button(actions, text="Download Excel", command=self.export_excel, state="disabled")
                self.export_button.pack(side="left", expand=True)

        points_panel = tk.Frame(parent, bg="#ffffff")
        points_panel.pack(fill="both", expand=True, pady=(8, 0))
        top_row = tk.Frame(points_panel, bg="#ffffff")
        top_row.pack(fill="x")
        tk.Label(top_row, text="Data points", bg="#ffffff", fg="#19233a", font=("Segoe UI", 12, "bold")).pack(side="left")
        self.point_count_var = tk.StringVar(value="0 points")
        tk.Label(top_row, textvariable=self.point_count_var, bg="#ffffff", fg="#68738a", font=("Segoe UI", 10)).pack(side="right")

        table_frame = tk.Frame(points_panel, bg="#ffffff", bd=1, relief="solid")
        table_frame.pack(fill="both", expand=True, pady=(10, 0))

        columns = ("#", "X", "Y")
        self.points_tree = ttk.Treeview(table_frame, columns=columns, show="headings", height=10)
        for col in columns:
            self.points_tree.heading(col, text=col)
            self.points_tree.column(col, anchor="center", width=90 if col != "#" else 40)
        self.points_tree.pack(fill="both", expand=True)

        self.tree_scroll = ttk.Scrollbar(table_frame, orient="vertical", command=self.points_tree.yview)
        self.points_tree.configure(yscrollcommand=self.tree_scroll.set)
        self.tree_scroll.pack(side="right", fill="y")

        action_row = tk.Frame(points_panel, bg="#ffffff")
        action_row.pack(fill="x", pady=(8, 0))
        self.remove_button = ttk.Button(action_row, text="Remove selected", command=self.remove_selected_point, state="disabled")
        self.remove_button.pack(anchor="w")

    def _field(self, parent, label_text):
        frame = tk.Frame(parent, bg="#ffffff")
        frame.grid_rowconfigure(0, weight=1)
        frame.grid_columnconfigure(0, weight=1)
        label = tk.Label(frame, text=label_text, bg="#ffffff", fg="#59657c", font=("Segoe UI", 9, "bold"))
        label.pack(anchor="w", padx=(0, 0), pady=(0, 4))
        entry = tk.Entry(frame, width=16, justify="center")
        entry.pack(fill="x", padx=(0, 0))
        return entry

    def _on_axis_input(self):
        self.message_var.set("")
        self._update_points()

    def open_file_dialog(self):
        file_path = filedialog.askopenfilename(
            title="Choose graph image or PDF",
            filetypes=[
                ("Image files", "*.png *.jpg *.jpeg *.bmp *.gif *.tif *.tiff"),
                ("PDF files", "*.pdf"),
                ("All files", "*.*"),
            ],
        )
        if not file_path:
            return
        self.path_var.set(file_path)
        self.load_image(file_path)

    def _process_path_entry(self, _event=None):
        file_path = self.path_var.get().strip().strip('"')
        if file_path:
            self.load_image(file_path)

    def _render_pdf_page(self, pdf_path, page_index=0, plot_bounds=None):
        if fitz is None:
            raise RuntimeError("PyMuPDF is required to process PDF files. Install it with: python -m pip install pymupdf")
        with fitz.open(pdf_path) as document:
            if not document.page_count:
                raise ValueError("The PDF does not contain any pages.")
            if not 0 <= page_index < document.page_count:
                raise ValueError(f"PDF page {page_index + 1} does not exist.")
            page = document[page_index]
            clip = None
            if plot_bounds is not None:
                panel = fitz.Rect(plot_bounds)
                clip = fitz.Rect(
                    panel.x0 - 35,
                    panel.y0 - 35,
                    panel.x1 + 45,
                    panel.y1 + 45,
                ) & page.rect
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False, clip=clip)
            image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        return image

    def load_image(self, file_path):
        file_path = os.path.abspath(os.path.expanduser(file_path.strip().strip('"')))
        supported_extensions = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff", ".pdf"}
        if not os.path.isfile(file_path):
            self.status_var.set(f"File not found: {file_path}")
            if self.headless_mode:
                raise FileNotFoundError(f"Input file not found: {file_path}")
            return
        ext = os.path.splitext(file_path)[1].lower()
        if ext not in supported_extensions:
            self.status_var.set("Choose a supported image or PDF file.")
            return
        if Image is None or ImageTk is None:
            self.status_var.set(f"Pillow is not available to {sys.executable}.")
            if self.headless_mode:
                raise RuntimeError(f"Pillow is not available to {sys.executable}.")
            return
        selected_graph = None
        try:
            if ext == ".pdf":
                graph_candidates = _detect_pdf_graphs(file_path)
                if graph_candidates:
                    selected_graph = _select_pdf_graph(self.root, graph_candidates)
                    if selected_graph is None:
                        self.status_var.set("Graph selection cancelled; no output was created.")
                        return
                    image = self._render_pdf_page(
                        file_path,
                        page_index=selected_graph["page_index"],
                        plot_bounds=selected_graph["plot_bounds"],
                    )
                else:
                    image = self._render_pdf_page(file_path)
            else:
                with Image.open(file_path) as opened:
                    image = opened.convert("RGBA")
        except Exception as exc:
            self.status_var.set(f"Could not open file: {exc}")
            if self.headless_mode:
                raise RuntimeError(f"The input file could not be opened: {exc}") from exc
            return

        self.state["image"] = image
        self.state["image_path"] = file_path
        self.state["mode"] = None
        self.state["origin"] = None
        self.state["opposite"] = None
        self.state["points"] = []
        self.state["curves"] = []
        self.state["axis_scales"] = {"x": "linear", "y": "linear"}
        self.state["axis_values"] = {"xMin": 0.0, "xMax": 10.0, "yMin": 0.0, "yMax": 10.0}
        self.state["chart_title"] = (
            selected_graph["title"]
            if selected_graph
            else os.path.splitext(os.path.basename(file_path))[0]
        )
        self.state["axis_titles"] = {"x": "X", "y": "Y"}
        self.last_export_path = None
        self.range_drag_start = None
        self.range_drag_current = None

        self.status_var.set(f"Processing {os.path.basename(file_path)}…")
        self.root.update_idletasks()
        if selected_graph:
            curve_data = _extract_pdf_curves(
                file_path,
                page_index=selected_graph["page_index"],
                requested_plot_bounds=selected_graph["plot_bounds"],
            )
            if curve_data is None:
                curve_data = _extract_pdf_raster_curves(
                    file_path,
                    page_index=selected_graph["page_index"],
                    plot_bounds=selected_graph["plot_bounds"],
                    chart_title=selected_graph["title"],
                )
        else:
            curve_data = _extract_pdf_curves(file_path) if ext == ".pdf" else None
        if selected_graph and curve_data is None:
            self.status_var.set(
                f"Detected '{selected_graph['title']}', but its curves could not be extracted automatically. "
                "No Excel file was created."
            )
            return
        if curve_data is not None:
            self.state["curves"] = curve_data["curves"]
            self.state["origin"] = (curve_data["plot_bounds"][0], curve_data["plot_bounds"][3])
            self.state["opposite"] = (curve_data["plot_bounds"][2], curve_data["plot_bounds"][1])
            self.state["axis_scales"] = curve_data["axis_scales"]
            self.state["axis_values"] = curve_data["axis_values"]
            self.state["chart_title"] = (
                selected_graph["title"]
                if selected_graph
                else curve_data["chart_title"]
            )
            self.state["axis_titles"] = curve_data["axis_titles"]
            self.state["points"] = curve_data["curves"][0]["points"]
        else:
            self._auto_calibrate_from_image()
            self._auto_collect_points()
        output_path = self.export_excel(run_headless=True)
        if output_path:
            self.last_export_path = output_path
            curve_count = len(self.state["curves"])
            details = f"Detected {curve_count} curves. " if curve_count else ""
            self.status_var.set(f"{details}Excel data saved to: {output_path}")
        elif not self.state["curves"] and not self.state["points"]:
            self.status_var.set("No graph data could be detected automatically; no Excel file was created.")

    def _auto_calibrate_from_image(self):
        image = self.state["image"]
        if image is None:
            return

        self.state["axis_values"] = {"xMin": 0.0, "xMax": 10.0, "yMin": 0.0, "yMax": 10.0}

        bounds = self._detect_plot_bounds(image)
        if bounds is None:
            return

        left, bottom, right, top = bounds
        self.state["origin"] = (left, bottom)
        self.state["opposite"] = (right, top)
        self.state["mode"] = None

    def _detect_plot_bounds(self, image):
        if image.mode not in {"RGBA", "RGB", "L"}:
            image = image.convert("RGBA")

        grayscale = image.convert("L")
        pixels = grayscale.load()
        width, height = grayscale.size

        dark_column_count = [0] * width
        dark_row_count = [0] * height
        for y in range(height):
            for x in range(width):
                value = pixels[x, y]
                if value < 220:
                    dark_column_count[x] += 1
                    dark_row_count[y] += 1

        column_threshold = max(4, int(height * 0.01))
        row_threshold = max(4, int(width * 0.01))
        left = next((x for x, count in enumerate(dark_column_count) if count > column_threshold), None)
        right = next((x for x, count in enumerate(reversed(dark_column_count), start=1) if count > column_threshold), None)
        top = next((y for y, count in enumerate(dark_row_count) if count > row_threshold), None)
        bottom = next((y for y, count in enumerate(reversed(dark_row_count), start=1) if count > row_threshold), None)

        if left is None or right is None or top is None or bottom is None:
            return None

        left = max(0, left - max(4, width // 30))
        right = min(width - 1, width - right + max(4, width // 30))
        top = max(0, top - max(4, height // 30))
        bottom = min(height - 1, height - bottom + max(4, height // 30))

        if right <= left or bottom <= top:
            return None

        padding = 0.02
        left_ratio = max(0.0, min(1.0, (left / width) + padding))
        right_ratio = max(0.0, min(1.0, (right / width) - padding))
        top_ratio = max(0.0, min(1.0, (top / height) + padding))
        bottom_ratio = max(0.0, min(1.0, (bottom / height) - padding))
        if right_ratio <= left_ratio or bottom_ratio <= top_ratio:
            return None
        return (left_ratio, bottom_ratio, right_ratio, top_ratio)

    def _auto_collect_points(self):
        image = self.state["image"]
        if image is None or not self._is_calibrated():
            return

        plot_image = image.convert("RGBA")
        width, height = plot_image.size
        left, bottom, right, top = (
            self.state["origin"][0] * width,
            self.state["origin"][1] * height,
            self.state["opposite"][0] * width,
            self.state["opposite"][1] * height,
        )
        left, right = sorted((max(0, int(left)), min(width - 1, int(right))))
        top, bottom = sorted((max(0, int(top)), min(height - 1, int(bottom))))

        crop = plot_image.crop((left, top, right, bottom))
        grayscale = crop.convert("L")
        pixels = grayscale.load()
        visited = set()
        markers = []
        threshold = 180

        for y in range(grayscale.height):
            for x in range(grayscale.width):
                if (x, y) in visited:
                    continue
                value = pixels[x, y]
                if value >= threshold:
                    continue
                region = [(x, y)]
                visited.add((x, y))
                stack = [(x, y)]
                while stack:
                    cx, cy = stack.pop()
                    for nx, ny in ((cx - 1, cy), (cx + 1, cy), (cx, cy - 1), (cx, cy + 1)):
                        if 0 <= nx < grayscale.width and 0 <= ny < grayscale.height and (nx, ny) not in visited:
                            visited.add((nx, ny))
                            if pixels[nx, ny] < threshold:
                                region.append((nx, ny))
                                stack.append((nx, ny))

                if len(region) < 6 or len(region) > 2000:
                    continue
                xs = [pt[0] for pt in region]
                ys = [pt[1] for pt in region]
                if max(xs) - min(xs) > 0 and max(ys) - min(ys) > 0:
                    cx = sum(xs) / len(xs)
                    cy = sum(ys) / len(ys)
                    abs_x = (left + cx) / width
                    abs_y = (top + cy) / height
                    markers.append((abs_x, abs_y))

        filtered = []
        for marker in sorted(markers, key=lambda item: (item[0], item[1])):
            if not filtered or all(abs(marker[0] - existing[0]) > 0.01 or abs(marker[1] - existing[1]) > 0.01 for existing in filtered):
                filtered.append(marker)

        if filtered:
            self.state["points"] = filtered

    def set_range_mode(self):
        if self.state["image"] is None:
            return
        self.state["mode"] = "range"
        self.range_drag_start = None
        self.range_drag_current = None
        self._update_calibration_buttons()
        self.message_var.set("")
        self.canvas_hint.configure(text="Click and drag around the graph's plotting area.")
        self.canvas.configure(cursor="crosshair")

    def _start_canvas_action(self, event):
        if self.state["image"] is None:
            return
        if self.state["mode"] == "range":
            position = self._canvas_to_image_position(event.x, event.y)
            if position is None:
                return
            self.range_drag_start = position
            self.range_drag_current = position
            self.draw_graph()
            return
        if self._is_calibrated():
            position = self._canvas_to_image_position(event.x, event.y)
            if position is not None:
                self.state["points"].append(position)
                self._update_points()
                self.draw_graph()
        else:
            self.message_var.set("Select the graph range before adding data points.")

    def _update_canvas_action(self, event):
        if self.state["mode"] != "range" or self.range_drag_start is None:
            return
        position = self._canvas_to_image_position(event.x, event.y, clamp=True)
        self.range_drag_current = position
        self.draw_graph()

    def _finish_canvas_action(self, event):
        if self.state["mode"] != "range" or self.range_drag_start is None:
            return
        end = self._canvas_to_image_position(event.x, event.y, clamp=True)
        start = self.range_drag_start
        left, right = sorted((start[0], end[0]))
        top, bottom = sorted((start[1], end[1]))
        self.range_drag_start = None
        self.range_drag_current = None

        if right - left < 0.001 or bottom - top < 0.001:
            self.message_var.set("Drag a larger rectangle around the graph's plotting area.")
            self.draw_graph()
            return

        self.state["origin"] = (left, bottom)
        self.state["opposite"] = (right, top)
        self.state["mode"] = None
        self.canvas.configure(cursor="")
        self._update_calibration_buttons()
        self._update_points()
        self.canvas_hint.configure(text="Range selected. Click the graph to add data points.")
        self.message_var.set("")
        self.draw_graph()

    def _canvas_to_image_position(self, x, y, clamp=False):
        if self.display_size == (0, 0):
            return None
        image_x = self.canvas.canvasx(x)
        image_y = self.canvas.canvasy(y)
        width, height = self.display_size
        if clamp:
            image_x = min(width, max(0, image_x))
            image_y = min(height, max(0, image_y))
        elif not (0 <= image_x <= width and 0 <= image_y <= height):
            return None
        return (image_x / width, image_y / height)

    def _update_calibration_buttons(self):
        self.range_button.configure(
            text="Redraw graph range" if self._is_calibrated() else "Drag to select graph range",
            state="normal" if self.state["image"] is not None else "disabled",
        )
        if self.state["mode"] == "range":
            self.range_button.state(["pressed"])
        else:
            self.range_button.state(["!pressed"])

    def _read_axis_values(self):
        values = dict(self.state["axis_values"])
        if values["xMin"] == values["xMax"] or values["yMin"] == values["yMax"]:
            raise ValueError("The minimum and maximum for each axis must be different.")
        return values

    def _is_calibrated(self):
        return bool(self.state["origin"] and self.state["opposite"])

    def _on_mousewheel(self, event):
        factor = 1.25 if event.delta > 0 else 1 / 1.25
        self._zoom_at(event, factor)
        return "break"

    def _zoom_at(self, event, factor):
        if self.state["image"] is None:
            return
        self._set_zoom(self.zoom * factor, (event.x, event.y))

    def _change_zoom(self, factor):
        self._set_zoom(self.zoom * factor)

    def _fit_image(self):
        image = self.state["image"]
        if image is None:
            return
        available_width = max(1, self.canvas.winfo_width())
        available_height = max(1, self.canvas.winfo_height())
        if available_width <= 1:
            available_width = 700
        if available_height <= 1:
            available_height = 520
        fit = min(1.0, available_width / image.width, available_height / image.height)
        self._set_zoom(max(0.05, fit), reset_view=True)

    def _set_zoom(self, zoom, anchor=None, reset_view=False):
        if self.state["image"] is None:
            return
        new_zoom = min(8.0, max(0.05, zoom))
        if abs(new_zoom - self.zoom) < 1e-9:
            if reset_view:
                self.canvas.xview_moveto(0)
                self.canvas.yview_moveto(0)
            return

        if anchor is not None:
            anchor_canvas_x = self.canvas.canvasx(anchor[0])
            anchor_canvas_y = self.canvas.canvasy(anchor[1])
        else:
            anchor_canvas_x = anchor_canvas_y = None
        old_zoom = self.zoom
        self.zoom = new_zoom
        self.draw_graph()

        if reset_view:
            self.canvas.xview_moveto(0)
            self.canvas.yview_moveto(0)
        elif anchor is not None:
            width, height = self.display_size
            target_x = max(0, anchor_canvas_x * new_zoom / old_zoom - anchor[0])
            target_y = max(0, anchor_canvas_y * new_zoom / old_zoom - anchor[1])
            self.canvas.xview_moveto(target_x / max(1, width))
            self.canvas.yview_moveto(target_y / max(1, height))
        self.zoom_label.configure(text=f"{round(self.zoom * 100)}%")

    def coordinates(self, position):
        values = self._read_axis_values()
        origin_x, origin_y = self.state["origin"]
        opposite_x, opposite_y = self.state["opposite"]
        x_ratio = (position[0] - origin_x) / (opposite_x - origin_x)
        y_ratio = (origin_y - position[1]) / (origin_y - opposite_y)

        def interpolate(minimum, maximum, ratio, scale):
            if scale == "log":
                if minimum <= 0 or maximum <= 0:
                    raise ValueError("Logarithmic axes require positive minimum and maximum values.")
                return 10 ** (math.log10(minimum) + ratio * (math.log10(maximum) - math.log10(minimum)))
            return minimum + ratio * (maximum - minimum)

        return {
            "x": interpolate(values["xMin"], values["xMax"], x_ratio, self.state["axis_scales"]["x"]),
            "y": interpolate(values["yMin"], values["yMax"], y_ratio, self.state["axis_scales"]["y"]),
        }

    def _update_points(self):
        self.points_tree.delete(*self.points_tree.get_children())
        point_count = len(self.state["points"])
        if self.state["curves"]:
            self.point_count_var.set(f"{len(self.state['curves'])} curves")
        else:
            self.point_count_var.set(f"{point_count} {'point' if point_count == 1 else 'points'}")

        if point_count == 0:
            self.points_tree.insert("", "end", values=("", "", ""))
            self.points_tree.item(self.points_tree.get_children()[-1], text="No points yet")
            self.points_tree.set(self.points_tree.get_children()[-1], "#", "")
            self.points_tree.set(self.points_tree.get_children()[-1], "X", "")
            self.points_tree.set(self.points_tree.get_children()[-1], "Y", "")
        else:
            for index, point in enumerate(self.state["points"], start=1):
                if self._is_calibrated():
                    try:
                        result = self.coordinates(point)
                        x_str = self.number_format(result["x"])
                        y_str = self.number_format(result["y"])
                    except ValueError as exc:
                        self.message_var.set(str(exc))
                        x_str = "—"
                        y_str = "—"
                else:
                    x_str = "—"
                    y_str = "—"
                self.points_tree.insert("", "end", values=(str(index), x_str, y_str))

        has_points = point_count > 0
        self.undo_button.configure(state="normal" if has_points else "disabled")
        self.clear_button.configure(state="normal" if has_points else "disabled")
        self.export_button.configure(state="normal" if has_points and self._is_calibrated() else "disabled")
        self.remove_button.configure(state="normal" if has_points else "disabled")

    def undo_last_point(self):
        if self.state["points"]:
            self.state["points"].pop()
            self._update_points()
            self.draw_graph()

    def clear_points(self):
        self.state["points"] = []
        self._update_points()
        self.draw_graph()

    def remove_selected_point(self):
        selection = self.points_tree.selection()
        if not selection:
            return
        item_id = selection[0]
        index = int(self.points_tree.set(item_id, "#")) - 1
        if 0 <= index < len(self.state["points"]):
            del self.state["points"][index]
            self._update_points()
            self.draw_graph()

    def draw_graph(self):
        self.canvas.delete("all")
        image = self.state["image"]
        if image is None:
            self.display_size = (0, 0)
            self.canvas.configure(scrollregion=(0, 0, 1, 1))
            return
        if Image is None or ImageTk is None:
            self.display_size = (0, 0)
            self.canvas.configure(scrollregion=(0, 0, 1, 1))
            return

        image_width = max(1, round(image.width * self.zoom))
        image_height = max(1, round(image.height * self.zoom))
        resampling = getattr(Image, "Resampling", Image).LANCZOS
        resized = image.resize((image_width, image_height), resampling)
        self.display_photo = ImageTk.PhotoImage(resized)
        self.display_size = (image_width, image_height)
        self.zoom_label.configure(text=f"{round(self.zoom * 100)}%")
        self.canvas.create_image(0, 0, anchor="nw", image=self.display_photo)
        self.canvas.configure(scrollregion=(0, 0, image_width, image_height))

        if self.state["origin"] and self.state["opposite"]:
            left = self.state["origin"][0] * image_width
            top = self.state["opposite"][1] * image_height
            right = self.state["opposite"][0] * image_width
            bottom = self.state["origin"][1] * image_height
            self.canvas.create_rectangle(left, top, right, bottom, outline="#365cf5", dash=(6, 5), width=2)

        if self.range_drag_start and self.range_drag_current:
            x1 = self.range_drag_start[0] * image_width
            y1 = self.range_drag_start[1] * image_height
            x2 = self.range_drag_current[0] * image_width
            y2 = self.range_drag_current[1] * image_height
            self.canvas.create_rectangle(x1, y1, x2, y2, outline="#13a06f", dash=(5, 3), width=2)

        if self.state["origin"]:
            self._draw_marker(self.state["origin"], "#13a06f", "1")
        if self.state["opposite"]:
            self._draw_marker(self.state["opposite"], "#13a06f", "2")

        if not self.state["curves"]:
            for index, position in enumerate(self.state["points"], start=1):
                self._draw_marker(position, "#ef625d", str(index))

    def _draw_marker(self, position, color, label):
        x = position[0] * self.display_size[0]
        y = position[1] * self.display_size[1]
        radius = max(4, min(self.canvas.winfo_width(), self.canvas.winfo_height()) / 140)
        self.canvas.create_oval(x - radius, y - radius, x + radius, y + radius, fill=color, outline="white", width=2)
        self.canvas.create_text(x + radius + 4, y, text=label, fill=color, font=("Segoe UI", 11, "bold"), anchor="w")

    def _write_curves_excel(self, out_path):
        if (
            Workbook is None
            or Reference is None
            or ScatterChart is None
            or Series is None
            or ChartLines is None
            or GraphicalProperties is None
            or Layout is None
            or ManualLayout is None
        ):
            raise RuntimeError("openpyxl is required for Excel export. Install it with: python -m pip install openpyxl")
        workbook = Workbook()
        combined_sheet = workbook.active
        combined_sheet.title = "Combined"
        axis_values = self._read_axis_values()
        data_header_row = 52
        combined_curve_name_row = data_header_row
        combined_axis_header_row = data_header_row + 1
        combined_data_start_row = combined_axis_header_row + 2
        curve_data_start_row = data_header_row + 1
        curves = self.state["curves"] or [
            {"name": "Curve 1", "color": "#365cf5", "points": self.state["points"]}
        ]
        converted_curves = [
            [self.coordinates(point) for point in curve["points"]]
            for curve in curves
        ]

        combined_sheet["A1"] = "Combined chart with local curve data below."
        combined_sheet.column_dimensions["A"].width = 20

        for curve_index, (curve, points) in enumerate(zip(curves, converted_curves)):
            sheet = workbook.create_sheet(curve["name"])
            sheet.column_dimensions["A"].width = 20
            sheet.column_dimensions["B"].width = 20

            x_column = curve_index * 2 + 1
            y_column = x_column + 1
            combined_sheet.cell(
                row=combined_curve_name_row,
                column=x_column,
                value=curve["name"],
            )
            combined_sheet.cell(
                row=combined_axis_header_row,
                column=x_column,
                value=self.state["axis_titles"]["x"],
            )
            combined_sheet.cell(
                row=combined_axis_header_row,
                column=y_column,
                value=self.state["axis_titles"]["y"],
            )
            for row_index, point in enumerate(points, start=combined_data_start_row):
                combined_sheet.cell(row=row_index, column=x_column, value=point["x"])
                combined_sheet.cell(row=row_index, column=y_column, value=point["y"])
            sheet.cell(row=data_header_row, column=1, value=self.state["axis_titles"]["x"])
            sheet.cell(row=data_header_row, column=2, value=self.state["axis_titles"]["y"])
            for row_index, point in enumerate(points, start=curve_data_start_row):
                sheet.cell(row=row_index, column=1, value=point["x"])
                sheet.cell(row=row_index, column=2, value=point["y"])
            combined_sheet.column_dimensions[
                combined_sheet.cell(row=1, column=x_column).column_letter
            ].width = 18
            combined_sheet.column_dimensions[
                combined_sheet.cell(row=1, column=y_column).column_letter
            ].width = 18

        def create_chart(title):
            chart = ScatterChart()
            chart.title = title
            chart.title.overlay = False
            chart.style = 13
            chart.scatterStyle = "line"
            chart.x_axis.title = self.state["axis_titles"]["x"]
            chart.y_axis.title = self.state["axis_titles"]["y"]
            chart.x_axis.axPos = "b"
            chart.y_axis.axPos = "l"
            chart.x_axis.delete = False
            chart.y_axis.delete = False
            chart.x_axis.tickLblPos = "low"
            chart.y_axis.tickLblPos = "nextTo"
            chart.x_axis.majorTickMark = "out"
            chart.y_axis.majorTickMark = "out"
            chart.x_axis.numFmt = "0.##"
            chart.y_axis.numFmt = "0.##"
            chart.x_axis.minorGridlines = ChartLines()
            chart.y_axis.minorGridlines = ChartLines()
            chart.x_axis.crosses = "min"
            chart.y_axis.crosses = "min"
            if self.state["axis_scales"]["x"] == "log":
                chart.x_axis.scaling.logBase = 10
                chart.x_axis.majorUnit = 1
            if self.state["axis_scales"]["y"] == "log":
                chart.y_axis.scaling.logBase = 10
                chart.y_axis.majorUnit = 1
            chart.x_axis.scaling.min = min(axis_values["xMin"], axis_values["xMax"])
            chart.x_axis.scaling.max = max(axis_values["xMin"], axis_values["xMax"])
            chart.y_axis.scaling.min = min(axis_values["yMin"], axis_values["yMax"])
            chart.y_axis.scaling.max = max(axis_values["yMin"], axis_values["yMax"])
            chart.plot_area.spPr = GraphicalProperties(solidFill="E8E8E8")
            chart.layout = Layout(
                manualLayout=ManualLayout(
                    layoutTarget="outer",
                    xMode="factor",
                    yMode="factor",
                    wMode="factor",
                    hMode="factor",
                    x=0.06,
                    y=0.04,
                    w=0.86,
                    h=0.84,
                )
            )
            chart.legend.position = "r"
            chart.legend.overlay = False
            chart.width = 22
            chart.height = 24
            return chart

        chart = create_chart(self.state["chart_title"])
        for curve_index, (curve, points) in enumerate(zip(curves, converted_curves)):
            x_column = curve_index * 2 + 1
            y_column = x_column + 1
            x_values = Reference(
                combined_sheet,
                min_col=x_column,
                min_row=combined_data_start_row,
                max_row=combined_data_start_row + len(points) - 1,
            )
            y_values = Reference(
                combined_sheet,
                min_col=y_column,
                min_row=combined_data_start_row,
                max_row=combined_data_start_row + len(points) - 1,
            )
            series = Series(y_values, x_values, title=curve["name"])
            series.smooth = False
            series.marker.symbol = "none"
            series.graphicalProperties.line.solidFill = curve["color"].lstrip("#")
            # Scale the stroke for the larger exported plot so source overlaps remain visible.
            series.graphicalProperties.line.width = 28575
            chart.series.append(series)
        combined_sheet.add_chart(chart, "A3")

        for curve_index, (curve, sheet) in enumerate(zip(curves, workbook.worksheets[1:])):
            chart = create_chart(f"{self.state['chart_title']} - {curve['name']}")
            chart.legend = None
            x_values = Reference(
                sheet,
                min_col=1,
                min_row=curve_data_start_row,
                max_row=curve_data_start_row + len(converted_curves[curve_index]) - 1,
            )
            y_values = Reference(
                sheet,
                min_col=2,
                min_row=curve_data_start_row,
                max_row=curve_data_start_row + len(converted_curves[curve_index]) - 1,
            )
            series = Series(y_values, x_values, title=curve["name"])
            series.smooth = False
            series.marker.symbol = "none"
            series.graphicalProperties.line.solidFill = curve["color"].lstrip("#")
            series.graphicalProperties.line.width = 28575
            chart.series.append(series)
            sheet.add_chart(chart, "A3")
        workbook.save(out_path)

    def export_excel(self, run_headless=False):
        if (not self.state["points"] and not self.state["curves"]) or not self._is_calibrated():
            return
        try:
            if run_headless:
                out_dir = os.path.dirname(self.state["image_path"]) or os.getcwd()
                basename = os.path.splitext(os.path.basename(self.state["image_path"]))[0]
                out_path = os.path.join(out_dir, f"{basename}-digitized.xlsx")
            else:
                out_path = filedialog.asksaveasfilename(
                    defaultextension=".xlsx",
                    filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
                    initialfile="digitized-data.xlsx",
                )
                if not out_path:
                    return
            self._write_curves_excel(out_path)
            self.last_export_path = out_path
            self.status_var.set(f"Excel data saved to: {out_path}")
            if not run_headless:
                messagebox.showinfo("Export complete", f"Saved Excel data to:\n{out_path}")
            return out_path
        except ValueError as exc:
            self.status_var.set(str(exc))
            return None
        except (OSError, RuntimeError) as exc:
            self.status_var.set(f"Excel export failed: {exc}")
            if not run_headless:
                messagebox.showerror("Excel export failed", str(exc))
            return None

def main():
    headless_path = len(sys.argv) > 1
    root = tk.Tk()
    if headless_path:
        root.withdraw()
    try:
        app = GraphDigitizerApp(root)
        if headless_path:
            if app.last_export_path:
                print(f"Saved Excel output: {app.last_export_path}")
                return 0
            error = app.status_var.get() or "No graph data could be detected; no Excel file was created."
            print(error, file=sys.stderr)
            return 1
        root.mainloop()
        return 0
    except (FileNotFoundError, OSError, RuntimeError, ValueError) as exc:
        print(str(exc), file=sys.stderr)
        return 2
    finally:
        if root.winfo_exists():
            root.destroy()


if __name__ == "__main__":
    raise SystemExit(main())
